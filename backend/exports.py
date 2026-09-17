import io
import re
from datetime import datetime
from pathlib import Path
from typing import Callable, List, Optional, Tuple

from docx import Document
from docx.shared import Inches, Pt
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image as RLImage, ListFlowable, ListItem, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ASSETS = Path(__file__).parent / "assets"
IMG_RE = re.compile(r"!\[([^\]]*)\]\(([^)\s]+)\)")
INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
HR_RE = re.compile(r"\s*(-{3,}|\*{3,}|_{3,})\s*")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)")
BULLET_RE = re.compile(r"^\s*(?:[-*+]|\d+[.)])\s+(.*)")
EMOJI_RE = re.compile("[\U0001F000-\U0001FAFF\U00002600-\U000027BF\U0001F900-\U0001F9FF\uFE0F\u200d]")
SEO_FIELDS = (("Meta Title", "meta_title"), ("Meta Description", "meta_description"), ("URL Slug", "url_slug"))
PLATFORM_LABELS = {"facebook": "Facebook", "instagram": "Instagram", "linkedin": "LinkedIn", "twitter": "X / Twitter", "tiktok": "TikTok"}
PLATFORM_COLORS = {"facebook": "#1877F2", "instagram": "#E1306C", "linkedin": "#0A66C2", "twitter": "#111111", "tiktok": "#FE2C55"}

ImageFetcher = Callable[[str], Optional[bytes]]
Block = Tuple[str, object]


# ---------- markdown parsing ----------

class _BlockCollector:
    def __init__(self):
        self.blocks: List[Block] = []
        self.para: List[str] = []
        self.items: List[str] = []

    def flush(self):
        if self.para:
            self.blocks.append(("p", " ".join(self.para)))
            self.para = []
        if self.items:
            self.blocks.append(("ul", self.items))
            self.items = []

    def add(self, kind: str, value) -> None:
        self.flush()
        self.blocks.append((kind, value))


def _classify(line: str) -> Tuple[str, object]:
    stripped = line.strip()
    if not stripped:
        return "blank", None
    if HR_RE.fullmatch(line):
        return "hr", None
    if img := IMG_RE.fullmatch(stripped):
        return "img", (img.group(1), img.group(2))
    if heading := HEADING_RE.match(line):
        return f"h{len(heading.group(1))}", heading.group(2).strip()
    if bullet := BULLET_RE.match(line):
        return "li", bullet.group(1)
    return "text", stripped


def parse_blocks(markdown: str) -> List[Block]:
    c = _BlockCollector()
    for raw in markdown.splitlines():
        kind, value = _classify(raw.rstrip())
        if kind == "blank":
            c.flush()
        elif kind == "li":
            if c.para:
                c.flush()
            c.items.append(value)
        elif kind == "text":
            if c.items:
                c.flush()
            c.para.append(value)
        else:
            c.add(kind, value)
    c.flush()
    return c.blocks


def inline_runs(text: str) -> List[Tuple[str, dict]]:
    runs = []
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            runs.append((part[2:-2], {"bold": True}))
        elif part.startswith("*"):
            runs.append((part[1:-1], {"italic": True}))
        elif part.startswith("`"):
            runs.append((part[1:-1], {"code": True}))
        elif part.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            runs.append((m.group(1), {"link": m.group(2)}))
        else:
            runs.append((part, {}))
    return runs


def _scaled(image_bytes: bytes, max_w: float, max_h: Optional[float] = None) -> Tuple[float, float]:
    with PILImage.open(io.BytesIO(image_bytes)) as im:
        w, h = im.size
    scale = min(1.0, max_w / w)
    if max_h:
        scale = min(scale, max_h / h)
    return w * scale, h * scale


def _strip_emoji(text: str) -> str:
    return re.sub(r"  +", " ", EMOJI_RE.sub("", text)).strip()


def _escape(text: str) -> str:
    return _strip_emoji(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# ---------- DOCX ----------

def _docx_runs(paragraph, text: str) -> None:
    for run_text, fmt in inline_runs(text):
        run = paragraph.add_run(run_text)
        run.bold = fmt.get("bold", False)
        run.italic = fmt.get("italic", False)
        if fmt.get("code"):
            run.font.name = "Courier New"


def _docx_image(doc, value, fetch_image: ImageFetcher) -> None:
    alt, url = value
    data = fetch_image(url)
    if not data:
        return
    doc.add_picture(io.BytesIO(data), width=Inches(6))
    if alt:
        cap = doc.add_paragraph(alt)
        cap.runs[0].italic = True
        cap.runs[0].font.size = Pt(9)


def _is_heading(kind: str) -> bool:
    return len(kind) == 2 and kind[0] == "h" and kind[1].isdigit()


def _docx_block(doc, kind: str, value, fetch_image: ImageFetcher) -> None:
    if _is_heading(kind):
        doc.add_heading(value, level=min(int(kind[1]), 4))
    elif kind == "p":
        _docx_runs(doc.add_paragraph(), value)
    elif kind == "ul":
        for item in value:
            _docx_runs(doc.add_paragraph(style="List Bullet"), item)
    elif kind == "hr":
        doc.add_paragraph()
    elif kind == "img":
        _docx_image(doc, value, fetch_image)


def build_docx(article: dict, fetch_image: ImageFetcher) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Georgia"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(article["title"], level=0)
    for kind, value in parse_blocks(article["content"]):
        _docx_block(doc, kind, value, fetch_image)

    doc.add_page_break()
    doc.add_heading("SEO Metadata", level=1)
    for label, key in SEO_FIELDS:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(article.get(key, ""))

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ---------- PDF (shared) ----------

def _register_fonts():
    if "Serif" in pdfmetrics.getRegisteredFontNames():
        return
    pdfmetrics.registerFont(TTFont("Serif", str(ASSETS / "LiberationSerif-Regular.ttf")))
    pdfmetrics.registerFont(TTFont("Serif-Bold", str(ASSETS / "LiberationSerif-Bold.ttf")))
    pdfmetrics.registerFont(TTFont("Serif-Italic", str(ASSETS / "LiberationSerif-Italic.ttf")))
    pdfmetrics.registerFont(TTFont("Sans-Bold", str(ASSETS / "LiberationSans-Bold.ttf")))
    pdfmetrics.registerFontFamily("Serif", normal="Serif", bold="Serif-Bold", italic="Serif-Italic", boldItalic="Serif-Bold")


def _inline_markup(text: str) -> str:
    parts = []
    for run, fmt in inline_runs(text):
        s = _escape(run)
        if fmt.get("bold"):
            s = f"<b>{s}</b>"
        elif fmt.get("italic"):
            s = f"<i>{s}</i>"
        elif fmt.get("code"):
            s = f"<font face='Courier'>{s}</font>"
        elif fmt.get("link"):
            s = f"<a href='{_escape(fmt['link'])}' color='#1d4ed8'>{s}</a>"
        parts.append(s)
    return "".join(parts)


def _paragraphs(text: str, style) -> list:
    return [Paragraph(_escape(line) or "&nbsp;", style) for line in text.split("\n")]


def _article_styles() -> dict:
    return {
        "body": ParagraphStyle("body", fontName="Serif", fontSize=11, leading=17, spaceAfter=9, alignment=TA_LEFT),
        "title": ParagraphStyle("title", fontName="Serif-Bold", fontSize=26, leading=32, spaceAfter=18),
        "h1": ParagraphStyle("h1", fontName="Sans-Bold", fontSize=18, leading=23, spaceBefore=14, spaceAfter=8),
        "h2": ParagraphStyle("h2", fontName="Sans-Bold", fontSize=15, leading=20, spaceBefore=12, spaceAfter=6),
        "h3": ParagraphStyle("h3", fontName="Sans-Bold", fontSize=12.5, leading=17, spaceBefore=10, spaceAfter=4),
        "caption": ParagraphStyle("caption", fontName="Serif-Italic", fontSize=9, leading=12, textColor="#666666", spaceAfter=10),
        "meta": ParagraphStyle("meta", fontName="Serif", fontSize=10, leading=14, textColor="#444444", spaceAfter=4),
    }


def _pdf_image(value, fetch_image: ImageFetcher, max_w: float, styles: dict) -> list:
    alt, url = value
    data = fetch_image(url)
    if not data:
        return []
    w, h = _scaled(data, max_w)
    flow = [Spacer(1, 6), RLImage(io.BytesIO(data), width=w, height=h)]
    if alt:
        flow.append(Paragraph(_escape(alt), styles["caption"]))
    return flow


def _pdf_block(kind: str, value, fetch_image: ImageFetcher, max_w: float, styles: dict) -> list:
    if _is_heading(kind):
        return [Paragraph(_inline_markup(value), styles.get(kind, styles["h3"]))]
    if kind == "p":
        return [Paragraph(_inline_markup(value), styles["body"])]
    if kind == "ul":
        return [ListFlowable([ListItem(Paragraph(_inline_markup(i), styles["body"]), leftIndent=12) for i in value], bulletType="bullet", start="•", leftIndent=14)]
    if kind == "hr":
        return [Spacer(1, 12)]
    if kind == "img":
        return _pdf_image(value, fetch_image, max_w, styles)
    return []


def build_pdf(article: dict, fetch_image: ImageFetcher) -> bytes:
    _register_fonts()
    styles = _article_styles()
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=0.9 * inch, rightMargin=0.9 * inch, topMargin=0.9 * inch, bottomMargin=0.9 * inch,
                            title=article.get("meta_title") or article["title"], author="Content Studio")
    max_w = A4[0] - 1.8 * inch
    story = [Paragraph(_escape(article["title"]), styles["title"])]
    for kind, value in parse_blocks(article["content"]):
        story += _pdf_block(kind, value, fetch_image, max_w, styles)
    story += [Spacer(1, 24), Paragraph("SEO Metadata", styles["h2"])]
    for label, key in SEO_FIELDS:
        story.append(Paragraph(f"<b>{label}:</b> {_escape(article.get(key, ''))}", styles["meta"]))
    doc.build(story)
    return out.getvalue()


# ---------- Campaign deck ----------

def _deck_styles(accent: str) -> dict:
    return {
        "cover": ParagraphStyle("cover", fontName="Sans-Bold", fontSize=34, leading=40, textColor="#111111", spaceAfter=14),
        "sub": ParagraphStyle("sub", fontName="Serif", fontSize=14, leading=20, textColor="#444444", spaceAfter=6),
        "kicker": ParagraphStyle("kicker", fontName="Sans-Bold", fontSize=10, leading=13, textColor=accent, spaceAfter=4),
        "h": ParagraphStyle("h", fontName="Sans-Bold", fontSize=22, leading=27, textColor="#111111", spaceAfter=10),
        "body": ParagraphStyle("body", fontName="Serif", fontSize=11.5, leading=17, textColor="#222222", spaceAfter=8),
        "tags": ParagraphStyle("tags", fontName="Serif-Italic", fontSize=10, leading=14, textColor="#555555"),
        "small": ParagraphStyle("small", fontName="Serif", fontSize=9, leading=12, textColor="#777777"),
    }


def _fmt_when(iso: Optional[str]) -> str:
    if not iso:
        return "Not scheduled"
    return datetime.fromisoformat(iso.replace("Z", "+00:00")).strftime("%a %b %d, %Y · %H:%M UTC")


def _page_decorator(accent: str, footer_text: str, page):
    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(colors.HexColor(accent))
        canvas.rect(0, page[1] - 8, page[0], 8, stroke=0, fill=1)
        canvas.setFont("Serif", 8)
        canvas.setFillColor(colors.HexColor("#888888"))
        canvas.drawString(0.8 * inch, 0.4 * inch, footer_text)
        canvas.drawRightString(page[0] - 0.8 * inch, 0.4 * inch, f"Page {doc.page}")
        canvas.restoreState()
    return on_page


def _deck_cover(campaign: dict, brand: dict, fetch_image: ImageFetcher, st: dict) -> list:
    logo = fetch_image(f"/api/files/{brand['logo_path']}") if brand.get("logo_path") else None
    story = []
    if logo:
        lw, lh = _scaled(logo, 1.6 * inch)
        story += [RLImage(io.BytesIO(logo), width=lw, height=lh, hAlign="LEFT"), Spacer(1, 24)]
    else:
        story.append(Spacer(1, 60))
    story += [Paragraph("CAMPAIGN DECK · FOR APPROVAL", st["kicker"]), Paragraph(_escape(campaign["name"]), st["cover"]),
              Paragraph(f"<b>Topic:</b> {_escape(campaign['topic'])}", st["sub"])]
    if campaign.get("goal"):
        story.append(Paragraph(f"<b>Goal:</b> {_escape(campaign['goal'])}", st["sub"]))
    scheduled = sum(1 for p in campaign["posts"] if p.get("scheduled_at"))
    summary = f"{len(campaign['posts'])} posts · {scheduled} scheduled" + (" · includes email" if campaign.get("email_copy") else "")
    return story + [Paragraph(summary, st["small"]), PageBreak()]


def _deck_post_page(post: dict, fetch_image: ImageFetcher, st: dict, content_w: float, max_h: float, accent: str) -> list:
    label = PLATFORM_LABELS.get(post["platform"], post["platform"].title())
    color = PLATFORM_COLORS.get(post["platform"], accent)
    left = [Paragraph(f"<font color='{color}'>■</font> {label.upper()}", st["kicker"]), Paragraph(_escape(label + " post"), st["h"]),
            Paragraph(_fmt_when(post.get("scheduled_at")), st["small"]), Spacer(1, 10)] + _paragraphs(post["content"], st["body"])
    if post.get("hashtags"):
        left.append(Paragraph(_escape(post["hashtags"]), st["tags"]))

    col_w = content_w * 0.46
    graphic = fetch_image(f"/api/files/{post['graphic_path']}") if post.get("graphic_path") else None
    if graphic:
        gw, gh = _scaled(graphic, col_w, max_h)
        right = [RLImage(io.BytesIO(graphic), width=gw, height=gh)]
    else:
        right = [Paragraph("No graphic created yet", st["small"])]
    table = Table([[left, right]], colWidths=[content_w - col_w - 12, col_w], hAlign="LEFT")
    table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (0, 0), 12)]))
    return [table, PageBreak()]


def _deck_schedule(campaign: dict, st: dict, content_w: float, accent: str) -> list:
    rows = [["When", "Platform", "Post"]]
    for p in sorted(campaign["posts"], key=lambda x: x.get("scheduled_at") or "9999"):
        excerpt = p["content"][:140] + ("…" if len(p["content"]) > 140 else "")
        rows.append([Paragraph(_fmt_when(p.get("scheduled_at")), st["body"]), Paragraph(PLATFORM_LABELS.get(p["platform"], p["platform"]), st["body"]),
                     Paragraph(_escape(excerpt), st["body"])])
    table = Table(rows, colWidths=[content_w * 0.26, content_w * 0.16, content_w * 0.58], hAlign="LEFT", repeatRows=1)
    table.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, 0), "Sans-Bold"), ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#555555")),
                               ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.HexColor(accent)), ("LINEBELOW", (0, 1), (-1, -1), 0.3, colors.HexColor("#DDDDDD")),
                               ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    return [Paragraph("SCHEDULE", st["kicker"]), Paragraph("Publishing schedule", st["h"]), Spacer(1, 8), table]


def build_campaign_deck(campaign: dict, brand: dict, fetch_image: ImageFetcher) -> bytes:
    _register_fonts()
    accent = brand.get("accent_color") or "#F59E0B"
    st = _deck_styles(accent)
    page = landscape(A4)
    content_w = page[0] - 1.6 * inch
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=page, leftMargin=0.8 * inch, rightMargin=0.8 * inch, topMargin=0.7 * inch, bottomMargin=0.7 * inch,
                            title=f"{campaign['name']} — Campaign Deck", author=brand.get("name") or "Content Studio")

    story = _deck_cover(campaign, brand, fetch_image, st)
    for post in campaign["posts"]:
        story += _deck_post_page(post, fetch_image, st, content_w, page[1] - 2.2 * inch, accent)
    if campaign.get("email_copy"):
        story += [Paragraph("EMAIL", st["kicker"]), Paragraph("Email blast", st["h"]), Spacer(1, 8)] + _paragraphs(campaign["email_copy"], st["body"]) + [PageBreak()]
    story += _deck_schedule(campaign, st, content_w, accent)

    footer = " · ".join(filter(None, [brand.get("name"), brand.get("handle"), campaign["name"]]))
    on_page = _page_decorator(accent, footer, page)
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return out.getvalue()
