"""Iteration 10 regression: exports refactor + graphics.zip."""
import io
import os
import zipfile

import fitz  # pymupdf
import pytest
import requests
from docx import Document
from PIL import Image

BASE_URL = os.environ["REACT_APP_BACKEND_URL"].rstrip("/")
ARTICLE_ID = "aa74874f-d686-459b-8bcd-0e36dc8fe3e2"
CAMPAIGN_ID = "12e4723b-5378-46cd-9d5c-42b5d63b77c3"


# ---------- article exports ----------

def test_article_pdf_download_has_images():
    r = requests.get(f"{BASE_URL}/api/articles/{ARTICLE_ID}/download/pdf", timeout=90)
    assert r.status_code == 200
    assert r.content[:4] == b"%PDF"
    doc = fitz.open(stream=r.content, filetype="pdf")
    total_images = sum(len(doc[i].get_images(full=True)) for i in range(doc.page_count))
    assert total_images >= 1, f"expected embedded images, got {total_images}"
    doc.close()


def test_article_docx_download_valid_with_headings_bullets_image():
    r = requests.get(f"{BASE_URL}/api/articles/{ARTICLE_ID}/download/docx", timeout=90)
    assert r.status_code == 200
    doc = Document(io.BytesIO(r.content))
    styles = [p.style.name for p in doc.paragraphs]
    texts = [p.text for p in doc.paragraphs]
    assert any(s.startswith("Heading") for s in styles), f"no headings in styles: {set(styles)}"
    assert any(s == "List Bullet" for s in styles), f"no list bullets in {set(styles)}"
    # `---` HR should NOT appear as text
    assert not any(t.strip() == "---" for t in texts), "raw --- present in docx"
    # inline picture check
    doc_xml = doc.part.blob.decode("utf-8", errors="ignore")
    images = [p for p in doc.part.related_parts.values() if "image" in getattr(p, "content_type", "")]
    assert len(images) >= 1, f"no inline images embedded, related={len(images)}, xml has drawing={'w:drawing' in doc_xml}"


def test_article_xlsx_download_400():
    r = requests.get(f"{BASE_URL}/api/articles/{ARTICLE_ID}/download/xlsx", timeout=30)
    assert r.status_code == 400


# ---------- campaign deck ----------

def test_campaign_deck_landscape_and_pages():
    c = requests.get(f"{BASE_URL}/api/campaigns/{CAMPAIGN_ID}", timeout=30).json()
    posts = len(c["posts"])
    email = 1 if c.get("email_copy") else 0
    expected = 1 + posts + email + 1

    r = requests.get(f"{BASE_URL}/api/campaigns/{CAMPAIGN_ID}/deck.pdf", timeout=90)
    assert r.status_code == 200
    assert r.content[:4] == b"%PDF"
    doc = fitz.open(stream=r.content, filetype="pdf")
    assert doc.page_count == expected, f"expected {expected} pages, got {doc.page_count}"
    w, h = doc[0].rect.width, doc[0].rect.height
    assert w > h, f"deck should be landscape, got {w}x{h}"
    doc.close()


def test_campaign_deck_unknown_404():
    r = requests.get(f"{BASE_URL}/api/campaigns/does-not-exist/deck.pdf", timeout=30)
    assert r.status_code == 404


# ---------- graphics.zip ----------

def test_graphics_zip_ok_content_and_filenames():
    c = requests.get(f"{BASE_URL}/api/campaigns/{CAMPAIGN_ID}", timeout=30).json()
    graphic_platforms = [p["platform"] for p in c["posts"] if p.get("graphic_path")]
    assert graphic_platforms, "seed campaign should have graphics"

    r = requests.get(f"{BASE_URL}/api/campaigns/{CAMPAIGN_ID}/graphics.zip", timeout=60)
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/zip"
    cd = r.headers.get("content-disposition", "")
    assert "-graphics.zip" in cd

    with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
        names = zf.namelist()
        assert names, "zip empty"
        # slug/<platform>.png structure
        for platform in graphic_platforms:
            match = [n for n in names if n.endswith(f"/{platform}.png")]
            assert match, f"missing {platform}.png in {names}"
            data = zf.read(match[0])
            img = Image.open(io.BytesIO(data))
            img.verify()
            assert img.format == "PNG"


def test_graphics_zip_unknown_campaign_404():
    r = requests.get(f"{BASE_URL}/api/campaigns/does-not-exist/graphics.zip", timeout=30)
    assert r.status_code == 404
